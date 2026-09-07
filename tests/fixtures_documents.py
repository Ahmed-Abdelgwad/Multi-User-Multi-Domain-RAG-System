"""Real PDF/DOCX byte fixtures for ingestion extraction tests -- avoids
mocking pypdf/python-docx/LangChain/Camelot themselves, so the tests
exercise the actual parsing libraries against real files.
"""
from io import BytesIO
import docx as python_docx
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet


def make_pdf_bytes(text: str | None = "Hello World") -> bytes:
    """A minimal single-page PDF. `text=None` produces a blank page (no
    extractable native text), useful for exercising the OCR-fallback
    trigger condition.
    """
    buf = BytesIO()
    c = canvas.Canvas(buf)
    if text:
        c.drawString(100, 750, text)
    c.save()
    return buf.getvalue()


def make_pdf_with_malformed_creation_date_bytes(text: str = "Hello World") -> bytes:
    """A real PDF whose `/CreationDate` is a raw string pypdf's own
    `meta.creation_date` property fails to parse -- regression fixture
    for a bug found live (not by any prior fixture) against a real
    resume PDF: `D:20260415123806Z'` (a trailing `'` with no UTC-offset
    minutes) raised inside `PdfReader(...).metadata.creation_date`,
    crashing extraction entirely even though the text itself was fine.
    """
    reader = PdfReader(BytesIO(make_pdf_bytes(text)))
    writer = PdfWriter()
    writer.append(reader)
    writer.add_metadata({"/CreationDate": "D:20260415123806Z'"})
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_pdf_with_paragraphs_bytes(paragraphs: list[str]) -> bytes:
    """A real multi-paragraph PDF via reportlab `Paragraph` flowables --
    used to prove `extraction._reconstruct_paragraphs` recovers real
    paragraph boundaries from PyPDFLoader's line-based output, which (per
    live verification) has no blank lines between paragraphs at all,
    unlike Docx2txtLoader.
    """
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    doc.build([Paragraph(p, styles["Normal"]) for p in paragraphs])
    return buf.getvalue()


def make_pdf_with_table_bytes(intro_text: str = "A document with a table") -> bytes:
    """A real PDF with a bordered (lattice-detectable) table, for Camelot
    extraction tests -- same shape used to validate Camelot manually
    against this project's Docker image.
    """
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    data = [
        ["Name", "Role", "Domain"],
        ["Alice", "Domain Admin", "Finance"],
        ["Bob", "Contributor", "HR"],
    ]
    table = Table(data)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
    ]))
    styles = getSampleStyleSheet()
    doc.build([Paragraph(intro_text, styles["Title"]), table])
    return buf.getvalue()


def make_docx_bytes(text: str = "Hello World", author: str | None = None) -> bytes:
    document = python_docx.Document()
    if author:
        document.core_properties.author = author
    document.add_paragraph(text)
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_docx_with_table_bytes(intro_text: str = "A document with a table") -> bytes:
    document = python_docx.Document()
    document.add_paragraph(intro_text)
    table = document.add_table(rows=3, cols=3)
    data = [
        ["Name", "Role", "Domain"],
        ["Alice", "Domain Admin", "Finance"],
        ["Bob", "Contributor", "HR"],
    ]
    for row_idx, row_data in enumerate(data):
        for col_idx, value in enumerate(row_data):
            table.cell(row_idx, col_idx).text = value
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()
